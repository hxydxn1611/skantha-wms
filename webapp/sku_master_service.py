"""
AI Vision & SKU Master Data Service
Phase 3: Excel Master Data Import, Validation, Template Generation,
and Filename-Based Product Image Mapping.
"""

import os
import re
import io
import logging
from typing import Dict, List, Any, Tuple, Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

try:
    import webapp.db as db
except ImportError:
    import db

logger = logging.getLogger(__name__)

# Allowed image extensions & max file size (10 MB)
ALLOWED_IMAGE_EXTENSIONS = {'.jpg', '.jpeg', '.png', '.webp', '.bmp'}
MAX_IMAGE_FILE_SIZE = 10 * 1024 * 1024

# Column synonym mappings for Excel header matching
COLUMN_SYNONYMS = {
    "sku_id": ["sku", "sku code", "sku_code", "sku id", "skuid", "item code", "item_code", "product code", "product_code"],
    "name": ["product name", "product_name", "item name", "item_name", "name", "title", "product title"],
    "description": ["description", "desc", "details", "product description", "item description"],
    "category": ["category", "cat", "product category", "group", "product group", "type"],
    "uom": ["uom", "unit of measure", "unit_of_measure", "unit", "units", "measure"],
    "units_per_bundle": [
        "individuals per bundle", "individuals_per_bundle", "units per bundle", "units_per_bundle",
        "pieces per bundle", "pieces_per_bundle", "pcs per bundle", "items per bundle", "bundle qty", "pack size"
    ],
    "bundle_name": ["bundle / pack type", "bundle type", "bundle_name", "pack type", "packaging type", "bundle"],
    "target_bundles_per_pallet": ["target bundles / pallet", "target bundles", "bundles per pallet", "target_bundles_per_pallet"],
    "unit_weight": ["unit weight", "unit_weight", "weight (kg)", "weight", "unit weight (kg)"],
    "unit_volume": ["unit volume", "unit_volume", "volume (m3)", "volume", "unit volume (m3)"],
    "status": ["status", "active status", "item status"]
}


def generate_sku_excel_template() -> bytes:
    """
    Generates a professional .xlsx master data template with clear required columns,
    practical industrial UOMs, manual Individuals per Bundle declaration, and styling.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    if ws is None:
        ws = wb.create_sheet()
    ws.title = "SKU_Master_Template"

    headers = [
        "SKU Code *",
        "Product Name *",
        "Description",
        "Category",
        "UOM *",
        "Individuals per Bundle *",
        "Bundle / Pack Type",
        "Target Bundles / Pallet",
        "Status"
    ]

    # Header styling (Corporate Navy with Bold White Text)
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    ws.append(headers)
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
    ws.row_dimensions[1].height = 28

    # Sample rows for user guidance across Glass, Wood/Timber, Metal, and Packaging
    sample_data = [
        ["GLS-FLT-012", "12mm Clear Toughened Float Glass", "Architectural jumbo tempered float glass 3660x2440mm", "Float Glass", "SHEETS", 25, "A-Frame Stillage", 2, "Active"],
        ["WOD-PLY-018", "18mm BWP Marine Grade Plywood", "Boiling waterproof hardwood structural plywood 8x4ft", "Plywood", "SHEETS", 40, "Strapped Plywood Stack", 4, "Active"],
        ["WOD-TMB-004", "Seasoned Burma Teak Timber Plank", "Kiln-dried structural timber planks 2x6 inch x 10ft", "Timber & Lumber", "PLANKS", 30, "Banded Timber Pack", 4, "Active"],
        ["MTL-PIP-050", "50mm Galvanized Steel Conduit Pipe", "Heavy gauge hot-dip galvanized structural tubes 6m", "Metals & Tubes", "PIPES", 45, "Hexagonal Steel Bundle", 6, "Active"],
        ["SKU-PKG-002", "Heavy Duty Corrugated Master Carton", "Double-wall corrugated shipping cartons 40x30x25cm", "Packaging", "BOXES", 20, "Pallet Layer Pack", 8, "Active"],
    ]

    sample_font = Font(name="Calibri", size=10)
    for row_idx, row_values in enumerate(sample_data, start=2):
        ws.append(row_values)
        for col_idx in range(1, len(row_values) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = sample_font
            cell.border = thin_border
            if col_idx in (1, 5, 9):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx in (6, 8):
                cell.alignment = Alignment(horizontal="right", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[row_idx].height = 20

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column or 1)
        for cell in col:
            val_str = str(cell.value or '')
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 5, 14)

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def _normalize_header_name(raw_name: Any) -> Optional[str]:
    """Matches a spreadsheet column title to a canonical SKU attribute name."""
    if not raw_name:
        return None
    cleaned = str(raw_name).strip().lower().replace("*", "").strip()
    for canonical, synonyms in COLUMN_SYNONYMS.items():
        if cleaned == canonical or cleaned in synonyms:
            return canonical
    return None


def parse_and_validate_sku_excel(file_bytes: bytes, duplicate_policy: str = "skip") -> Dict[str, Any]:
    """
    Parses an uploaded .xlsx workbook, performs strict schema & row-level validation,
    checks existing database records, and generates a structured preview.
    """
    try:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
        ws = wb.active
        if ws is None:
            return {
                "success": False,
                "message": "The uploaded workbook contains no active worksheet.",
                "can_import": False,
                "rows": []
            }
    except Exception as e:
        return {
            "success": False,
            "message": f"Unable to parse Excel file: {str(e)}. Please upload a valid .xlsx spreadsheet.",
            "can_import": False,
            "rows": []
        }

    # Extract header row
    header_row = None
    header_col_map: Dict[int, str] = {}
    for row in ws.iter_rows(values_only=True):
        if any(cell is not None and str(cell).strip() != '' for cell in row):
            header_row = row
            break

    if not header_row:
        return {
            "success": False,
            "message": "The uploaded spreadsheet is empty.",
            "can_import": False,
            "rows": []
        }

    for col_idx, col_val in enumerate(header_row):
        canon = _normalize_header_name(col_val)
        if canon:
            header_col_map[col_idx] = canon

    if "sku_id" not in header_col_map.values():
        return {
            "success": False,
            "message": "Missing required column 'SKU Code' (or 'Item Code'). Please check the template.",
            "can_import": False,
            "rows": []
        }
    if "name" not in header_col_map.values():
        return {
            "success": False,
            "message": "Missing required column 'Product Name' (or 'Item Name'). Please check the template.",
            "can_import": False,
            "rows": []
        }

    # Retrieve all existing SKU codes from DB
    conn = db.get_db()
    existing_rows = conn.execute("SELECT sku_id FROM skus").fetchall()
    conn.close()
    existing_sku_set = {r['sku_id'].upper() for r in existing_rows}

    preview_rows = []
    seen_in_file = set()
    valid_new_count = 0
    valid_update_count = 0
    skipped_count = 0
    invalid_count = 0

    # Iterate rows starting after header
    row_num = 1
    passed_header = False
    for row in ws.iter_rows(values_only=True):
        row_num += 1
        if not passed_header:
            if row == header_row:
                passed_header = True
            continue

        # Skip completely empty rows
        if not any(cell is not None and str(cell).strip() != '' for cell in row):
            continue

        raw_item: Dict[str, Any] = {}
        for col_idx, val in enumerate(row):
            if col_idx in header_col_map:
                field = header_col_map[col_idx]
                raw_item[field] = val

        sku_code = str(raw_item.get("sku_id") or "").strip().upper()
        prod_name = str(raw_item.get("name") or "").strip()
        desc = str(raw_item.get("description") or "").strip()
        category = str(raw_item.get("category") or "General").strip() or "General"
        uom = str(raw_item.get("uom") or "EA").strip().upper() or "EA"
        bundle_name = str(raw_item.get("bundle_name") or "Bundle").strip() or "Bundle"
        status_val = str(raw_item.get("status") or "Active").strip().capitalize()
        if status_val not in ("Active", "Disabled"):
            status_val = "Active"

        # Validate numbers
        upb_err = None
        weight_err = None
        vol_err = None
        raw_upb = raw_item.get("units_per_bundle")
        raw_tbp = raw_item.get("target_bundles_per_pallet")
        raw_wt = raw_item.get("unit_weight")
        raw_vol = raw_item.get("unit_volume")

        try:
            units_per_bundle = int(float(raw_upb)) if raw_upb is not None and str(raw_upb).strip() != '' else 1
            if units_per_bundle < 1:
                upb_err = "Individuals per Bundle must be >= 1"
                units_per_bundle = 1
        except (ValueError, TypeError):
            upb_err = f"Invalid Individuals per Bundle '{raw_upb}'"
            units_per_bundle = 1

        try:
            target_bundles_per_pallet = max(1, int(float(raw_tbp))) if raw_tbp is not None and str(raw_tbp).strip() != '' else 10
        except (ValueError, TypeError):
            target_bundles_per_pallet = 10

        try:
            unit_weight = float(raw_wt) if raw_wt is not None and str(raw_wt).strip() != '' else 0.0
            if unit_weight < 0:
                weight_err = "Weight cannot be negative"
        except (ValueError, TypeError):
            weight_err = f"Invalid weight value '{raw_wt}'"
            unit_weight = 0.0

        try:
            unit_volume = float(raw_vol) if raw_vol is not None and str(raw_vol).strip() != '' else 0.0
            if unit_volume < 0:
                vol_err = "Volume cannot be negative"
        except (ValueError, TypeError):
            vol_err = f"Invalid volume value '{raw_vol}'"
            unit_volume = 0.0

        row_errors = []
        if not sku_code:
            row_errors.append("SKU Code is required.")
        if not prod_name:
            row_errors.append("Product Name is required.")
        if upb_err:
            row_errors.append(upb_err)
        if weight_err:
            row_errors.append(weight_err)
        if vol_err:
            row_errors.append(vol_err)

        # Duplicate within file
        if sku_code:
            if sku_code in seen_in_file:
                row_errors.append(f"Duplicate SKU '{sku_code}' within the uploaded file.")
            else:
                seen_in_file.add(sku_code)

        # Determine proposed action
        if row_errors:
            action = "REJECT"
            action_label = "Rejected"
            invalid_count += 1
        elif sku_code in existing_sku_set:
            if duplicate_policy == "update":
                action = "UPDATE"
                action_label = "Will Update"
                valid_update_count += 1
            else:
                action = "SKIP"
                action_label = "Will Skip (Exists)"
                skipped_count += 1
        else:
            action = "INSERT"
            action_label = "Will Insert"
            valid_new_count += 1

        preview_rows.append({
            "row_num": row_num,
            "sku_id": sku_code,
            "name": prod_name,
            "description": desc,
            "category": category,
            "uom": uom,
            "units_per_bundle": units_per_bundle,
            "bundle_name": bundle_name,
            "target_bundles_per_pallet": target_bundles_per_pallet,
            "unit_weight": round(unit_weight, 4),
            "unit_volume": round(unit_volume, 4),
            "status": status_val,
            "action": action,
            "action_label": action_label,
            "error": "; ".join(row_errors) if row_errors else None
        })

    total_valid = valid_new_count + (valid_update_count if duplicate_policy == 'update' else 0)
    return {
        "success": True,
        "total_rows": len(preview_rows),
        "valid_new_count": valid_new_count,
        "valid_update_count": valid_update_count,
        "skipped_count": skipped_count,
        "invalid_count": invalid_count,
        "can_import": total_valid > 0,
        "rows": preview_rows
    }


def match_bulk_images_to_skus(filenames: List[str]) -> List[Dict[str, Any]]:
    """
    Performs automatic filename-based matching to associate image files with SKU records.
    Pattern: 'SKU001.jpg', 'SKU-001_front.png', 'SKU_001.webp' -> 'SKU-001' or 'SKU001'.
    Returns proposed mappings, matching confidence, and allows client-side correction.
    """
    conn = db.get_db()
    sku_rows = conn.execute("SELECT sku_id, name FROM skus").fetchall()
    conn.close()

    known_skus = {r['sku_id'].upper(): r['name'] for r in sku_rows}
    # Normalized lookup without hyphens/underscores for fuzzy tolerance
    canonical_lookup = {}
    for s_code, s_name in known_skus.items():
        clean_key = re.sub(r'[^A-Z0-9]', '', s_code)
        canonical_lookup[clean_key] = (s_code, s_name)

    results = []
    for fname in filenames:
        base = os.path.basename(fname)
        root, ext = os.path.splitext(base)
        ext_clean = ext.lower()

        is_valid_ext = ext_clean in ALLOWED_IMAGE_EXTENSIONS
        if not is_valid_ext:
            results.append({
                "filename": base,
                "matched_sku_id": None,
                "matched_name": None,
                "status": "INVALID_TYPE",
                "status_label": "Invalid File Type",
                "error": f"Unsupported format '{ext}'. Allowed: {', '.join(ALLOWED_IMAGE_EXTENSIONS)}"
            })
            continue

        # Strategy 1: Exact root match
        candidate_code = root.strip().upper()
        if candidate_code in known_skus:
            results.append({
                "filename": base,
                "matched_sku_id": candidate_code,
                "matched_name": known_skus[candidate_code],
                "status": "MATCHED",
                "status_label": "Exact Match",
                "error": None
            })
            continue

        # Strategy 2: Prefix before delimiter (e.g. 'TEST-SKU-MATCH_front.png', 'SKU-001_1.jpg', 'SKU-001-back.png')
        matched_prefix_sku = None
        for delim in ('_', '-', ' '):
            parts = root.rsplit(delim, 1)
            if len(parts) == 2 and parts[0].strip().upper() in known_skus:
                matched_prefix_sku = parts[0].strip().upper()
                break

        if not matched_prefix_sku:
            for s_code in sorted(known_skus.keys(), key=len, reverse=True):
                if candidate_code.startswith(s_code) and len(candidate_code) > len(s_code):
                    if candidate_code[len(s_code)] in ('_', '-', ' ', '.'):
                        matched_prefix_sku = s_code
                        break

        if matched_prefix_sku:
            results.append({
                "filename": base,
                "matched_sku_id": matched_prefix_sku,
                "matched_name": known_skus[matched_prefix_sku],
                "status": "MATCHED",
                "status_label": "Prefix Match",
                "error": None
            })
            continue

        # Strategy 3: Alpha-numeric key match (e.g. 'SKU001' matches 'SKU-001')
        clean_root = re.sub(r'[^A-Z0-9]', '', root.upper())
        if clean_root in canonical_lookup:
            matched_code, matched_name = canonical_lookup[clean_root]
            results.append({
                "filename": base,
                "matched_sku_id": matched_code,
                "matched_name": matched_name,
                "status": "MATCHED",
                "status_label": "Fuzzy Key Match",
                "error": None
            })
            continue

        # No match found
        results.append({
            "filename": base,
            "matched_sku_id": None,
            "matched_name": None,
            "status": "UNMATCHED",
            "status_label": "Unmatched SKU",
            "error": "No matching SKU found in Master Data. Please select manually."
        })

    return results
